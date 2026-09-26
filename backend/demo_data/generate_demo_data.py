import os
import time
from datetime import datetime
from pathlib import Path
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image, ImageDraw, ImageFont

def generate_demo_files(target_dir: str):
    folder = Path(target_dir)
    folder.mkdir(parents=True, exist_ok=True)

    # Specific timestamp: February 14, 2026, 14:30:00
    feb_dt = datetime(2026, 2, 14, 14, 30, 0)
    feb_timestamp = feb_dt.timestamp()

    # 1. final2.pdf (Target Document for main scenario)
    pdf_path = folder / "final2.pdf"
    c = canvas.Canvas(str(pdf_path), pagesize=letter)
    
    # Page 1
    c.setFont("Helvetica-Bold", 18)
    c.drawString(50, 750, "IoT Asset Tracking Platform: Architecture Specification")
    c.setFont("Helvetica", 10)
    c.drawString(50, 730, "Document ID: ARCH-2026-V3 | Author: Rahul Sharma (rahul@example.com)")
    c.drawString(50, 715, "Date: February 14, 2026 | Sent to: Core Engineering Team")
    c.line(50, 705, 550, 705)

    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, 680, "1. Executive Architecture Summary")
    c.setFont("Helvetica", 10)
    summary_text = (
        "This architecture document outlines the next-generation indoor positioning and telemetry system. "
        "The infrastructure relies on low-power BLE (Bluetooth Low Energy) beacon hardware communicating with "
        "field gateways, streaming real-time coordinate data to our distributed backend services. "
        "Primary persistence and spatial indexing are handled by a high-availability PostgreSQL cluster."
    )
    # Simple line wrapping
    words = summary_text.split()
    y = 660
    line = []
    for w in words:
        line.append(w)
        if len(" ".join(line)) > 75:
            c.drawString(50, y, " ".join(line[:-1]))
            line = [w]
            y -= 15
    if line:
        c.drawString(50, y, " ".join(line))
        y -= 25

    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, y, "2. Bluetooth Low Energy (BLE) Triangulation Protocol")
    y -= 20
    c.setFont("Helvetica", 10)
    ble_text = (
        "Field telemetry nodes broadcast iBeacon/Eddystone frames over BLE channels 37, 38, and 39. "
        "Gateway receivers measure RSSI values to compute trilateration coordinates within a 1.2m margin. "
        "BLE packets are packaged into protobuf envelopes before egress."
    )
    for w in ble_text.split():
        line.append(w)
        if len(" ".join(line)) > 75:
            c.drawString(50, y, " ".join(line[:-1]))
            line = [w]
            y -= 15
    if line:
        c.drawString(50, y, " ".join(line))
        y -= 25

    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, y, "3. Persistence Layer: PostgreSQL & PostGIS")
    y -= 20
    c.setFont("Helvetica", 10)
    pg_text = (
        "PostgreSQL serves as the system of record. Beacon telemetry is partitioned across relational tables "
        "optimized for time-series ingestion. PostgreSQL connection pooling is handled via PgBouncer. "
        "All asset historical paths are indexed in PostgreSQL using GiST indexing."
    )
    line = []
    for w in pg_text.split():
        line.append(w)
        if len(" ".join(line)) > 75:
            c.drawString(50, y, " ".join(line[:-1]))
            line = [w]
            y -= 15
    if line:
        c.drawString(50, y, " ".join(line))

    c.showPage()

    # Page 2
    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, 750, "Section 4: Backend Microservices & Ingestion Flow")
    c.setFont("Helvetica", 10)
    p2_text = (
        "Rahul's review notes: Ensure the BLE message queue maintains sub-50ms latency. "
        "The PostgreSQL database must have automated failover enabled. "
        "Sent by Rahul for technical architecture verification."
    )
    c.drawString(50, 720, p2_text)
    c.showPage()
    c.save()

    # Set modified and access time to Feb 14, 2026
    os.utime(str(pdf_path), (feb_timestamp, feb_timestamp))

    # 2. IMG_2384.png (Architecture Screenshot)
    img_path = folder / "IMG_2384.png"
    img = Image.new("RGB", (900, 600), color=(24, 24, 27))
    draw = ImageDraw.Draw(img)
    
    # Draw boxes representing architecture components
    # Header
    draw.text((40, 30), "SYSTEM ARCHITECTURE: IOT & BACKEND INFRASTRUCTURE", fill=(244, 244, 245))
    draw.text((40, 55), "Author: Rahul | Project: Core IoT Gateway", fill=(161, 161, 170))
    
    # BLE Nodes box
    draw.rectangle([60, 120, 260, 220], outline=(99, 102, 241), width=2, fill=(39, 39, 42))
    draw.text((80, 145), "BLE Beacons", fill=(255, 255, 255))
    draw.text((80, 175), "RSSI Telemetry Mesh", fill=(161, 161, 170))

    # Gateway box
    draw.rectangle([340, 120, 540, 220], outline=(99, 102, 241), width=2, fill=(39, 39, 42))
    draw.text((360, 145), "Gateway Services", fill=(255, 255, 255))
    draw.text((360, 175), "FastAPI / Python", fill=(161, 161, 170))

    # PostgreSQL box
    draw.rectangle([620, 120, 840, 220], outline=(16, 185, 129), width=2, fill=(39, 39, 42))
    draw.text((640, 145), "PostgreSQL Database", fill=(255, 255, 255))
    draw.text((640, 175), "pgvector + Spatial", fill=(161, 161, 170))

    # Connection lines
    draw.line([260, 170, 340, 170], fill=(244, 244, 245), width=2)
    draw.line([540, 170, 620, 170], fill=(244, 244, 245), width=2)

    # Footer note
    draw.text((60, 320), "Architecture notes: BLE packets processed in real-time and persisted to PostgreSQL.", fill=(212, 212, 216))
    img.save(str(img_path))
    os.utime(str(img_path), (feb_timestamp, feb_timestamp))

    # Companion metadata for IMG_2384.png
    meta_img_path = folder / "IMG_2384.png.meta.json"
    with open(meta_img_path, "w", encoding="utf-8") as f:
        f.write('''{
  "detected_text": [
    "SYSTEM ARCHITECTURE: IOT & BACKEND INFRASTRUCTURE",
    "Author: Rahul",
    "BLE Beacons",
    "RSSI Telemetry Mesh",
    "Gateway Services",
    "FastAPI / Python",
    "PostgreSQL Database",
    "pgvector + Spatial",
    "Architecture notes: BLE packets processed in real-time and persisted to PostgreSQL."
  ],
  "visual_description": "Software system architecture diagram illustrating BLE beacon receivers communicating with gateway services and storing spatial coordinate records in a PostgreSQL database.",
  "metadata": {
    "author": "Rahul",
    "date": "February 2026",
    "category": "architecture_diagram"
  }
}''')

    # 3. document.pdf (Marketing Plan)
    doc_pdf = folder / "document.pdf"
    c_m = canvas.Canvas(str(doc_pdf), pagesize=letter)
    c_m.setFont("Helvetica-Bold", 16)
    c_m.drawString(50, 750, "Q1 Growth Marketing & Acquisition Strategy")
    c_m.setFont("Helvetica", 10)
    c_m.drawString(50, 730, "Author: Sarah Jenkins | April 2026")
    c_m.drawString(50, 700, "This report evaluates digital marketing channels, pay-per-click spend, CAC, and LTV metrics.")
    c_m.drawString(50, 680, "No hardware engineering topics are covered in this budget review.")
    c_m.showPage()
    c_m.save()

    # 4. architecture_old.pdf (Legacy System 2024 with MySQL & Zigbee)
    arch_old = folder / "architecture_old.pdf"
    c_old = canvas.Canvas(str(arch_old), pagesize=letter)
    c_old.setFont("Helvetica-Bold", 16)
    c_old.drawString(50, 750, "Legacy Monolith Architecture Overview (2024 Archive)")
    c_old.setFont("Helvetica", 10)
    c_old.drawString(50, 730, "Author: David Miller | Published: September 2024")
    c_old.drawString(50, 700, "Historical document outlining our previous monolithic stack.")
    c_old.drawString(50, 680, "Database: MySQL 8.0 cluster. Sensor networking: Zigbee protocol mesh (Deprecated).")
    c_old.drawString(50, 660, "Note: Superseded by the new BLE + PostgreSQL architecture designed in 2026.")
    c_old.showPage()
    c_old.save()

    # 5. notes_final.pdf (Sprint retro with Rahul)
    notes_pdf = folder / "notes_final.pdf"
    c_n = canvas.Canvas(str(notes_pdf), pagesize=letter)
    c_n.setFont("Helvetica-Bold", 16)
    c_n.drawString(50, 750, "Sprint Sync & Weekly Meeting Minutes")
    c_n.setFont("Helvetica", 10)
    c_n.drawString(50, 730, "Attendees: Rahul, Joyel, Ananya | Date: March 2026")
    c_n.drawString(50, 700, "Discussed Next.js dashboard UI polish, user profile forms, and sprint retro actions.")
    c_n.drawString(50, 680, "Frontend components are ready for staging testing.")
    c_n.showPage()
    c_n.save()

    # 6. presentation2.pptx (Product launch slides)
    pptx_path = folder / "presentation2.pptx"
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)
    
    # Slide 1
    blank_layout = prs.slide_layouts[6]
    slide1 = prs.slides.add_slide(blank_layout)
    txBox = slide1.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Commercial Sensor Series - Product Launch"
    p.font.size = Pt(28)
    p.font.bold = True
    p2 = tf.add_paragraph()
    p2.text = "Q2 Roadmap & Hardware Deliverables | Hardware Ops"
    p2.font.size = Pt(14)

    # Slide 2
    slide2 = prs.slides.add_slide(blank_layout)
    txBox2 = slide2.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3))
    tf2 = txBox2.text_frame
    p_s2 = tf2.paragraphs[0]
    p_s2.text = "Key Hardware Specifications"
    p_s2.font.size = Pt(22)
    p_s2.font.bold = True
    p_s2_desc = tf2.add_paragraph()
    p_s2_desc.text = "Battery endurance: 36 months under standard 1Hz beaconing rate. IP67 waterproof housing."

    prs.save(str(pptx_path))

    # 7. report_v1.docx (Security Audit)
    docx_path = folder / "report_v1.docx"
    doc = Document()
    doc.add_heading("Cloud Security & Compliance Audit Report", level=1)
    doc.add_paragraph("Auditor: Infosec Team | Date: January 2026")
    doc.add_heading("1. Scope", level=2)
    doc.add_paragraph("Assessment of VPC endpoints, IAM role boundaries, TLS encryption in transit, and secret rotation policies.")
    doc.save(str(docx_path))

    print(f"Generated {len(list(folder.glob('*')))} demo files in {folder}")

if __name__ == "__main__":
    import sys
    target = sys.argv[1] if len(sys.argv) > 1 else "./demo_data"
    generate_demo_files(target)
