import logging
from typing import Dict, Any, List
from pathlib import Path
from pptx import Presentation

logger = logging.getLogger("refind.processors.pptx")

def process_pptx(file_path: str) -> Dict[str, Any]:
    """
    Extracts text, slide notes, and per-slide chunks from a PPTX file.
    """
    path = Path(file_path)
    chunks: List[Dict[str, Any]] = []
    full_text_parts: List[str] = []
    metadata: Dict[str, Any] = {}

    try:
        prs = Presentation(str(path))
        
        # Read core properties if available
        if prs.core_properties:
            cp = prs.core_properties
            if cp.title:
                metadata["title"] = cp.title
            if cp.author:
                metadata["author"] = cp.author

        for idx, slide in enumerate(prs.slides):
            slide_num = idx + 1
            slide_texts: List[str] = []

            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_texts.append(text)

            # Slide notes
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    slide_texts.append(f"Notes: {notes_text}")

            if slide_texts:
                content = "\n".join(slide_texts)
                full_text_parts.append(f"[Slide {slide_num}]\n{content}")
                chunks.append({
                    "content": content,
                    "page_number": None,
                    "slide_number": slide_num
                })
    except Exception as e:
        logger.error(f"Error processing PPTX {file_path}: {e}")

    extracted_text = "\n\n".join(full_text_parts)
    return {
        "extracted_text": extracted_text,
        "chunks": chunks,
        "metadata": metadata
    }
